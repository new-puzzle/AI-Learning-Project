"""
GoalPath AI - Universal Goal Planning Platform
Main Streamlit Application
"""

import streamlit as st
import os
from dotenv import load_dotenv
from datetime import datetime
from utils.path_generator import LearningPathGenerator
from utils.auth import init_cookie_manager, check_password, render_login_screen, logout

# Load environment variables
load_dotenv()

# Initialize cookie manager for authentication
cookies = init_cookie_manager()

# Check authentication before showing app
if not check_password(cookies):
    render_login_screen(cookies)
    st.stop()

# Page configuration
st.set_page_config(
    page_title="GoalPath AI",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom CSS for better styling and mobile optimization
st.markdown("""
    <style>
    /* Desktop styles */
    .main-header {
        font-size: 3rem;
        font-weight: bold;
        color: #1E88E5;
        text-align: center;
        margin-bottom: 0.5rem;
    }
    .sub-header {
        font-size: 1.2rem;
        color: #666;
        text-align: center;
        margin-bottom: 2rem;
    }
    .topic-card {
        background-color: #f5f5f5;
        padding: 1.5rem;
        border-radius: 10px;
        border-left: 4px solid #1E88E5;
        margin-bottom: 1rem;
    }
    .completed-topic {
        background-color: #e8f5e9;
        border-left: 4px solid #4CAF50;
    }
    .stat-box {
        background-color: #f0f2f6;
        padding: 1rem;
        border-radius: 8px;
        text-align: center;
    }
    .resource-link {
        color: #1E88E5;
        text-decoration: none;
    }

    /* Mobile optimization */
    @media only screen and (max-width: 768px) {
        .main-header {
            font-size: 2rem;
            margin-bottom: 0.3rem;
        }
        .sub-header {
            font-size: 1rem;
            margin-bottom: 1rem;
        }
        .topic-card {
            padding: 1rem;
        }
        .stat-box {
            padding: 0.75rem;
            font-size: 0.9rem;
        }
        /* Make buttons touch-friendly */
        .stButton > button {
            min-height: 44px;
            font-size: 1rem;
            padding: 12px 20px;
        }
        /* Improve form input sizes */
        .stTextInput > div > div > input,
        .stTextArea > div > div > textarea,
        .stSelectbox > div > div > select {
            font-size: 16px !important;
            min-height: 44px;
        }
        /* Better spacing on mobile */
        .block-container {
            padding: 1rem 0.5rem;
        }
        /* Optimize file uploader for mobile */
        .stFileUploader {
            font-size: 0.9rem;
        }
        /* Make YouTube embeds responsive */
        iframe {
            max-width: 100%;
        }
    }

    /* Extra small devices */
    @media only screen and (max-width: 480px) {
        .main-header {
            font-size: 1.5rem;
        }
        .sub-header {
            font-size: 0.9rem;
        }
    }
    </style>
""", unsafe_allow_html=True)


def init_session_state():
    """Initialize session state variables"""
    if 'current_path_id' not in st.session_state:
        st.session_state.current_path_id = None
    if 'show_generator' not in st.session_state:
        st.session_state.show_generator = True
    if 'generated_path' not in st.session_state:
        st.session_state.generated_path = None


def check_api_key():
    """Check if API key is configured"""
    api_key = os.getenv("ANTHROPIC_API_KEY")
    if not api_key or api_key == "your_api_key_here":
        st.error("⚠️ Anthropic API Key not configured!")
        st.info("""
        **To use LearnPath AI:**
        1. Copy `.env.example` to `.env`
        2. Add your Anthropic API key to the `.env` file
        3. Get your API key from: https://console.anthropic.com/
        4. Restart the application
        """)
        return False
    return True


def render_header():
    """Render the application header"""
    st.markdown('<p class="main-header">🎯 GoalPath AI</p>', unsafe_allow_html=True)
    st.markdown('<p class="sub-header">AI-Powered Goal Planner for Learning, Career, Freelancing & More</p>', unsafe_allow_html=True)


def render_learning_path_generator(generator):
    """Render the goal plan generator form"""
    st.markdown("### 🚀 Create Your Goal Plan")

    # Goal type selector
    goal_type_options = {
        "📚 Learning & Skills": "learning",
        "💼 Career Transition": "career",
        "💰 Freelance & Business": "freelance",
        "🚀 Project Completion": "project",
        "🎯 Personal Achievement": "personal"
    }

    selected_type = st.selectbox(
        "Goal Type",
        list(goal_type_options.keys()),
        help="Select the type of goal you want to achieve"
    )
    goal_type = goal_type_options[selected_type]

    # Update placeholder based on goal type
    placeholders = {
        "learning": "e.g., Learn prompt engineering, Master Python, Understand blockchain",
        "career": "e.g., Get hired as AI engineer, Transition to product management, Land remote job",
        "freelance": "e.g., Get 5 Fiverr clients, Earn $1000/month freelancing, Launch consulting business",
        "project": "e.g., Build portfolio website, Launch mobile app, Write a book",
        "personal": "e.g., Run a 10K race, Learn guitar, Lose 20 pounds"
    }

    col1, col2 = st.columns([2, 1])

    with col1:
        goal = st.text_input(
            "What do you want to achieve?",
            placeholder=placeholders.get(goal_type, "Enter your goal"),
            help="Enter your goal in a clear, specific way"
        )

    with col2:
        timeframe = st.number_input(
            "Timeframe (days)",
            min_value=1,
            max_value=365,
            value=30,
            help="How many days do you want to dedicate to this goal?"
        )

    if st.button("Generate Goal Plan", type="primary", use_container_width=True):
        if not goal:
            st.error("Please enter a goal!")
            return

        with st.spinner(f"🤖 AI is creating your personalized {selected_type.split()[0].lower()} plan..."):
            try:
                # Generate goal plan
                learning_path = generator.create_learning_path(goal, timeframe, goal_type)

                # Store in session state
                st.session_state.generated_path = learning_path
                st.session_state.current_path_id = learning_path['path_id']
                st.session_state.show_generator = False

                st.success(f"✅ {selected_type} plan generated successfully!")
                st.rerun()

            except Exception as e:
                st.error(f"Error generating goal plan: {str(e)}")


def render_learning_path(path_data):
    """Render the complete learning path"""
    st.markdown("### 📚 Your Learning Path")

    # Overview section
    if 'overview' in path_data:
        st.info(f"**Overview:** {path_data['overview']}")

    # Milestones
    if 'milestones' in path_data and path_data['milestones']:
        st.markdown("#### 🎯 Key Milestones")
        cols = st.columns(len(path_data['milestones']))
        for idx, milestone in enumerate(path_data['milestones']):
            with cols[idx]:
                st.markdown(f"**{idx + 1}.** {milestone}")

    # Curriculum
    st.markdown("#### 📖 Daily Curriculum")

    curriculum = path_data.get('curriculum', [])

    for topic in curriculum:
        day_num = topic.get('day', 0)
        topic_name = topic.get('topic', '')
        subtopics = topic.get('subtopics', [])
        estimated_hours = topic.get('estimated_hours', 0)
        resources = topic.get('resources', [])

        with st.expander(f"📅 Day {day_num}: {topic_name}", expanded=(day_num == 1)):
            col1, col2 = st.columns([3, 1])

            with col1:
                st.markdown(f"**Topic:** {topic_name}")

                if subtopics:
                    st.markdown("**Subtopics:**")
                    for subtopic in subtopics:
                        st.markdown(f"- {subtopic}")

            with col2:
                st.metric("Estimated Time", f"{estimated_hours} hours")

            # Resources
            if resources:
                st.markdown("**📚 Recommended Resources:**")
                for resource in resources:
                    resource_type = resource.get('type', 'resource')
                    resource_name = resource.get('name', 'Resource')
                    resource_url = resource.get('url', '#')

                    icon = {
                        'course': '🎓',
                        'article': '📄',
                        'video': '🎥',
                        'book': '📖'
                    }.get(resource_type, '🔗')

                    st.markdown(f"{icon} [{resource_name}]({resource_url})")


def render_saved_paths(generator):
    """Render list of saved learning paths"""
    st.markdown("### 📋 My Learning Paths")

    # Status filter
    status_filter = st.selectbox(
        "Filter by status:",
        ["All", "Active", "On Hold", "Archived"],
        key="status_filter"
    )

    # Map filter to database values
    status_map = {
        "All": None,
        "Active": "active",
        "On Hold": "on_hold",
        "Archived": "archived"
    }

    paths = generator.get_paths_by_status(status_map[status_filter])

    if not paths:
        st.info("No learning paths yet. Create your first one above!")
        return

    for path in paths:
        path_id = path['id']
        goal = path['goal']
        timeframe = path['timeframe']
        created_at = path['created_at']
        status = path.get('status', 'active')
        goal_type = path.get('goal_type', 'learning')

        # Get progress stats
        stats = generator.get_progress_stats(path_id)
        progress = stats['progress_percentage']

        # Status badge
        status_colors = {
            'active': '🟢',
            'on_hold': '🟡',
            'archived': '⚪',
            'deleted': '🔴'
        }
        status_icon = status_colors.get(status, '🟢')

        # Goal type icons
        goal_type_icons = {
            'learning': '📚',
            'career': '💼',
            'freelance': '💰',
            'project': '🚀',
            'personal': '🎯'
        }
        goal_type_icon = goal_type_icons.get(goal_type, '📚')

        col1, col2, col3 = st.columns([3, 1, 1])

        with col1:
            st.markdown(f"{status_icon} {goal_type_icon} **{goal}**")
            st.caption(f"Created: {created_at[:10]} | {timeframe} days | Type: {goal_type.title()} | Status: {status.replace('_', ' ').title()}")

        with col2:
            st.progress(progress / 100)
            st.caption(f"{progress:.0f}% complete")

        with col3:
            if st.button("View", key=f"view_{path_id}"):
                st.session_state.current_path_id = path_id
                st.session_state.show_generator = False
                st.rerun()


def render_progress_tracker(generator, path_id):
    """Render progress tracking interface"""
    path_data = generator.get_learning_path(path_id)

    if not path_data:
        st.error("Learning path not found!")
        return

    path_info = path_data['path_info']
    curriculum = path_data['curriculum']
    stats = path_data['stats']

    # Header with back button and status controls
    col1, col2, col3 = st.columns([3, 1, 1])
    with col1:
        st.markdown(f"### 📚 {path_info['goal']}")
    with col2:
        # Status management
        current_status = path_info.get('status', 'active')
        new_status = st.selectbox(
            "Status",
            ["active", "on_hold", "archived"],
            index=["active", "on_hold", "archived"].index(current_status),
            key=f"status_select_{path_id}",
            label_visibility="collapsed"
        )
        if new_status != current_status:
            generator.update_path_status(path_id, new_status)
            st.success(f"Status updated to: {new_status.replace('_', ' ').title()}")
            st.rerun()
    with col3:
        if st.button("← Back"):
            st.session_state.show_generator = True
            st.session_state.current_path_id = None
            st.rerun()

    # Export and advanced options
    with st.expander("⚙️ Advanced Options"):
        # PDF Export
        st.markdown("**📄 Export**")
        if st.button("📥 Download as PDF", key=f"export_pdf_{path_id}"):
            try:
                from io import BytesIO
                from reportlab.lib.pagesizes import letter
                from reportlab.lib.styles import getSampleStyleSheet
                from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
                from reportlab.lib.units import inch

                buffer = BytesIO()
                doc = SimpleDocTemplate(buffer, pagesize=letter)
                styles = getSampleStyleSheet()
                story = []

                # Title
                goal_type = path_info.get('goal_type', 'learning').title()
                title = Paragraph(f"<b>{path_info['goal']}</b>", styles['Title'])
                story.append(title)
                story.append(Spacer(1, 0.2*inch))

                # Metadata
                meta = Paragraph(f"Type: {goal_type} | Timeframe: {path_info['timeframe']} days | Progress: {stats['progress_percentage']:.0f}%", styles['Normal'])
                story.append(meta)
                story.append(Spacer(1, 0.3*inch))

                # Curriculum
                for topic in curriculum:
                    day_text = Paragraph(f"<b>Day {topic['day']}: {topic['topic']}</b>", styles['Heading2'])
                    story.append(day_text)
                    story.append(Spacer(1, 0.1*inch))

                    # Priority and hours
                    priority = topic.get('priority', 'medium').upper()
                    info = Paragraph(f"Priority: {priority} | Estimated: {topic['estimated_hours']} hours", styles['Normal'])
                    story.append(info)
                    story.append(Spacer(1, 0.1*inch))

                    # Subtopics
                    if topic['subtopics']:
                        for subtopic in topic['subtopics']:
                            bullet = Paragraph(f"• {subtopic}", styles['Normal'])
                            story.append(bullet)
                    story.append(Spacer(1, 0.2*inch))

                doc.build(story)
                buffer.seek(0)

                st.download_button(
                    label="💾 Save PDF",
                    data=buffer,
                    file_name=f"goalpath_{path_info['goal'][:30].replace(' ', '_')}.pdf",
                    mime="application/pdf"
                )
            except Exception as e:
                st.error(f"PDF generation failed: {str(e)}")

        st.markdown("---")
        st.warning("Danger Zone")
        if st.button("🗑️ Delete this learning path", key=f"delete_{path_id}"):
            if st.session_state.get(f"confirm_delete_{path_id}", False):
                generator.update_path_status(path_id, 'deleted')
                st.success("Learning path deleted!")
                st.session_state.show_generator = True
                st.session_state.current_path_id = None
                st.rerun()
            else:
                st.session_state[f"confirm_delete_{path_id}"] = True
                st.error("Click again to confirm deletion")

    # Progress stats
    col1, col2, col3, col4 = st.columns(4)

    with col1:
        st.markdown('<div class="stat-box">', unsafe_allow_html=True)
        st.metric("Progress", f"{stats['progress_percentage']:.0f}%")
        st.markdown('</div>', unsafe_allow_html=True)

    with col2:
        st.markdown('<div class="stat-box">', unsafe_allow_html=True)
        st.metric("Completed", f"{stats['completed_topics']}/{stats['total_topics']}")
        st.markdown('</div>', unsafe_allow_html=True)

    with col3:
        st.markdown('<div class="stat-box">', unsafe_allow_html=True)
        st.metric("Time Spent", f"{stats['total_time_spent_hours']} hrs")
        st.markdown('</div>', unsafe_allow_html=True)

    with col4:
        streak = generator.get_learning_streak(path_id)
        st.markdown('<div class="stat-box">', unsafe_allow_html=True)
        st.metric("Learning Days", streak)
        st.markdown('</div>', unsafe_allow_html=True)

    # Progress bar
    st.progress(stats['progress_percentage'] / 100)

    # On-track indicator
    from datetime import datetime
    created_at = path_info.get('created_at', '')
    timeframe = path_info.get('timeframe', 30)

    if created_at:
        try:
            # Parse created_at date (format: YYYY-MM-DD HH:MM:SS)
            created_date = datetime.strptime(created_at.split(' ')[0], '%Y-%m-%d')
            current_date = datetime.now()
            days_elapsed = (current_date - created_date).days

            # Calculate expected progress
            expected_progress = min((days_elapsed / timeframe) * 100, 100)
            actual_progress = stats['progress_percentage']

            # Determine status
            if actual_progress >= expected_progress:
                status_text = "✅ On Track"
                status_color = "green"
            else:
                behind_by = expected_progress - actual_progress
                status_text = f"⚠️ Behind by {behind_by:.0f}%"
                status_color = "orange"

            # Display on-track status
            col1, col2 = st.columns([1, 3])
            with col1:
                st.markdown(f"**Status:**")
            with col2:
                st.markdown(f"<span style='color: {status_color}; font-weight: bold;'>{status_text}</span> (Expected: {expected_progress:.0f}% by day {days_elapsed})", unsafe_allow_html=True)
        except:
            pass  # Skip if date parsing fails

    st.markdown("---")

    # Tabs for Curriculum and AI Tutor
    tab1, tab2 = st.tabs(["📖 Curriculum", "🤖 AI Tutor"])

    with tab1:
        # Curriculum with checkboxes
        st.markdown("#### 📖 Your Learning Curriculum")

        for topic in curriculum:
            topic_id = topic['id']
            day_num = topic['day']
            topic_name = topic['topic']
            subtopics = topic['subtopics']
            estimated_hours = topic['estimated_hours']
            resources = topic['resources']
            is_completed = topic['is_completed']
            priority = topic.get('priority', 'medium')

            # Priority indicators
            priority_emoji = {
                'high': '🔴',
                'medium': '🟡',
                'low': '🟢'
            }.get(priority, '🟡')

            # Topic card
            card_class = "completed-topic" if is_completed else ""

            with st.container():
                col1, col2 = st.columns([1, 20])

                with col1:
                    # Checkbox for completion
                    completed = st.checkbox(
                        f"Mark {topic_name} as complete",
                        value=is_completed,
                        key=f"topic_{topic_id}",
                        label_visibility="collapsed"
                    )

                    # Update if changed
                    if completed != is_completed:
                        generator.update_progress(topic_id, completed, 0)
                        st.rerun()

                with col2:
                    with st.expander(f"{priority_emoji} {'✅' if is_completed else '⭐'} Day {day_num}: {topic_name}", expanded=not is_completed):
                        st.markdown(f"**Estimated Time:** {estimated_hours} hours")

                        if subtopics:
                            st.markdown("**What you'll learn:**")
                            for subtopic in subtopics:
                                st.markdown(f"- {subtopic}")

                        if resources:
                            st.markdown("**📚 Resources:**")
                            for resource in resources:
                                resource_type = resource.get('type', 'resource')
                                resource_name = resource.get('name', 'Resource')
                                resource_url = resource.get('url', '#')

                                icon = {
                                    'course': '🎓',
                                    'article': '📄',
                                    'video': '🎥',
                                    'book': '📖'
                                }.get(resource_type, '🔗')

                                st.markdown(f"{icon} [{resource_name}]({resource_url})")

                                # YouTube auto-embed
                                if 'youtube.com' in resource_url or 'youtu.be' in resource_url:
                                    import re
                                    # Extract video ID
                                    video_id = None
                                    if 'youtube.com/watch?v=' in resource_url:
                                        video_id = resource_url.split('v=')[1].split('&')[0]
                                    elif 'youtu.be/' in resource_url:
                                        video_id = resource_url.split('youtu.be/')[1].split('?')[0]

                                    if video_id:
                                        st.markdown(f"""
                                        <iframe width="100%" height="315"
                                        src="https://www.youtube.com/embed/{video_id}"
                                        frameborder="0"
                                        allow="accelerometer; autoplay; clipboard-write; encrypted-media; gyroscope; picture-in-picture"
                                        allowfullscreen>
                                        </iframe>
                                        """, unsafe_allow_html=True)

            st.markdown("<br>", unsafe_allow_html=True)

    with tab2:
        # AI Tutor Chat Interface
        st.markdown("#### 🤖 AI Learning Tutor")
        st.caption("Ask questions, upload images/files, and get personalized explanations from multiple AI models")

        # Initialize chat history in session state
        if f'chat_history_{path_id}' not in st.session_state:
            st.session_state[f'chat_history_{path_id}'] = []

        # Get available models
        available_models = generator.get_available_models()

        # Build model options with provider grouping
        model_options = []
        for provider_name, provider_data in available_models.items():
            configured_status = "✓" if provider_data['configured'] else "⚠️"
            vision_status = "👁️" if provider_data['supports_vision'] else ""
            for model in provider_data['models']:
                model_options.append(f"{configured_status} {vision_status} {provider_name} / {model}")

        # Model selector
        st.markdown("**Select AI Model:** (✓ = configured, ⚠️ = not configured, 👁️ = supports images)")
        model_choice = st.selectbox(
            "Model",
            model_options,
            key=f"model_selector_{path_id}",
            label_visibility="collapsed",
            help="Configure API keys in .env file to enable all models"
        )

        # File upload
        st.markdown("**Upload Image or File (optional):**")
        uploaded_file = st.file_uploader(
            "Upload file",
            type=["png", "jpg", "jpeg", "pdf", "txt", "docx"],
            key=f"file_upload_{path_id}",
            label_visibility="collapsed",
            help="Upload images, PDFs, or documents for the AI to analyze"
        )

        # Context about current learning path
        current_context = f"Learning Goal: {path_info['goal']}\nTopics covered: " + ", ".join([t['topic'] for t in curriculum[:5]])

        # Display chat history
        for i, message in enumerate(st.session_state[f'chat_history_{path_id}']):
            if message['role'] == 'user':
                st.markdown(f"**You:** {message['content']}")
            else:
                model_used = message.get('model', 'Claude Sonnet 4.5')
                st.markdown(f"**AI Tutor ({model_used}):** {message['content']}")
            st.markdown("---")

        # Chat input
        user_question = st.text_area(
            "Ask a question about your learning topics:",
            placeholder="e.g., Can you explain prompt engineering in simple terms?\nWhat's the difference between zero-shot and few-shot prompting?",
            height=100,
            key=f"chat_input_{path_id}"
        )

        col1, col2 = st.columns([1, 5])
        with col1:
            if st.button("Ask Tutor", type="primary"):
                if user_question.strip() or uploaded_file:
                    # Clean up model choice (remove status icons)
                    clean_model = model_choice
                    for icon in ["✓", "⚠️", "👁️", " "]:
                        clean_model = clean_model.replace(icon, "")
                    clean_model = clean_model.strip()

                    # Prepare user message
                    user_msg = user_question if user_question.strip() else "(Analyzing uploaded file)"
                    if uploaded_file:
                        user_msg += f"\n📎 Uploaded: {uploaded_file.name}"

                    # Add user message to history
                    st.session_state[f'chat_history_{path_id}'].append({
                        'role': 'user',
                        'content': user_msg
                    })

                    # Get AI response
                    with st.spinner(f"AI Tutor ({clean_model}) is thinking..."):
                        try:
                            if uploaded_file and uploaded_file.type.startswith('image/'):
                                # Image analysis
                                image_data = uploaded_file.read()
                                prompt = user_question if user_question.strip() else "Analyze this image and explain what you see. If it contains any problems or questions, solve them."
                                response = generator.analyze_uploaded_image(image_data, prompt, clean_model)
                            elif uploaded_file:
                                # Text file handling (PDF, TXT, DOCX)
                                file_content = ""
                                if uploaded_file.type == "application/pdf":
                                    try:
                                        import PyPDF2
                                        import io
                                        pdf_reader = PyPDF2.PdfReader(io.BytesIO(uploaded_file.read()))
                                        for page in pdf_reader.pages:
                                            file_content += page.extract_text()
                                    except:
                                        file_content = "(Could not extract PDF text)"
                                elif uploaded_file.type == "text/plain":
                                    file_content = uploaded_file.read().decode('utf-8')
                                elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                                    try:
                                        import docx
                                        import io
                                        doc = docx.Document(io.BytesIO(uploaded_file.read()))
                                        file_content = "\n".join([para.text for para in doc.paragraphs])
                                    except:
                                        file_content = "(Could not extract document text)"

                                full_question = f"Based on this document:\n\n{file_content[:2000]}...\n\n{user_question}"
                                response = generator.get_assistance(full_question, current_context, clean_model)
                            else:
                                # Regular text question
                                response = generator.get_assistance(user_question, current_context, clean_model)

                            # Add AI response to history with model info
                            st.session_state[f'chat_history_{path_id}'].append({
                                'role': 'assistant',
                                'content': response,
                                'model': clean_model
                            })

                            st.rerun()
                        except Exception as e:
                            st.error(f"Error: {str(e)}\n\nMake sure the API key for this provider is configured in your .env file.")
                else:
                    st.warning("Please enter a question or upload a file")

        with col2:
            if st.button("Clear Chat History"):
                st.session_state[f'chat_history_{path_id}'] = []
                st.rerun()

        # Quick question suggestions
        st.markdown("**💡 Quick Question Suggestions:**")
        suggestions_col1, suggestions_col2 = st.columns(2)

        with suggestions_col1:
            if st.button("Explain this topic simply", key=f"suggest1_{path_id}"):
                st.session_state[f'suggested_question_{path_id}'] = "Can you explain the current topic in simple terms with examples?"

            if st.button("Give me practice problems", key=f"suggest2_{path_id}"):
                st.session_state[f'suggested_question_{path_id}'] = "Can you give me some practice problems for this topic?"

        with suggestions_col2:
            if st.button("What are common mistakes?", key=f"suggest3_{path_id}"):
                st.session_state[f'suggested_question_{path_id}'] = "What are common mistakes beginners make with this topic?"

            if st.button("How does this apply in real-world?", key=f"suggest4_{path_id}"):
                st.session_state[f'suggested_question_{path_id}'] = "How is this used in real-world applications?"


def main():
    """Main application"""
    init_session_state()
    render_header()

    # Check API key
    if not check_api_key():
        return

    try:
        # Initialize generator
        generator = LearningPathGenerator()

        # Sidebar
        with st.sidebar:
            st.markdown("### 🎯 Navigation")

            if st.button("➕ New Learning Path", use_container_width=True):
                st.session_state.show_generator = True
                st.session_state.current_path_id = None
                st.rerun()

            st.markdown("---")
            render_saved_paths(generator)

            st.markdown("---")
            st.markdown("### ℹ️ About")
            st.caption("GoalPath AI helps you create structured goal plans with AI-powered planning and progress tracking.")

            # Logout button
            st.markdown("---")
            if st.button("🚪 Logout", use_container_width=True, type="secondary"):
                logout(cookies)

        # Main content
        if st.session_state.current_path_id:
            # Show progress tracker
            render_progress_tracker(generator, st.session_state.current_path_id)
        elif st.session_state.show_generator:
            # Show generator form
            render_learning_path_generator(generator)

            # Show recently generated path
            if st.session_state.generated_path:
                st.markdown("---")
                render_learning_path(st.session_state.generated_path)
        else:
            st.info("Select a learning path from the sidebar or create a new one!")

    except Exception as e:
        st.error(f"Application error: {str(e)}")
        st.info("Please check your configuration and try again.")


if __name__ == "__main__":
    main()
